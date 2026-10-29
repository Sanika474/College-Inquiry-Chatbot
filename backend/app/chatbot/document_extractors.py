"""Document extraction utilities for CampusIQ Knowledge Base.

Supports extracting text, headings, tables, and page-by-page metadata from
PDFs, DOCX files, TXT, Markdown, CSV, and JSON files, with image/scanned PDF detection.
"""

import csv
import json
import os
from typing import Any, Dict, List, Optional
from pypdf import PdfReader
import docx


class ExtractedPage:
    """Represents a single page or section extracted from a document."""

    def __init__(
        self,
        page_number: int,
        text: str,
        section_heading: Optional[str] = None,
        tables: Optional[List[List[str]]] = None,
        is_empty: bool = False,
    ):
        self.page_number = page_number
        self.text = text.strip()
        self.section_heading = section_heading
        self.tables = tables or []
        self.is_empty = is_empty or (len(self.text) == 0 and not self.tables)


class DocumentExtractionResult:
    """Complete extraction result for a document file."""

    def __init__(
        self,
        filename: str,
        file_path: str,
        file_type: str,
        file_size_bytes: int,
        pages: List[ExtractedPage],
        status: str,  # 'success', 'scanned_ocr_required', 'unsupported', 'error'
        error_message: Optional[str] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ):
        self.filename = filename
        self.file_path = file_path
        self.file_type = file_type
        self.file_size_bytes = file_size_bytes
        self.pages = pages
        self.status = status
        self.error_message = error_message
        self.metadata = metadata or {}

    @property
    def total_pages(self) -> int:
        return len(self.pages)

    @property
    def total_characters(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)


class DocumentExtractor:
    """Extracts text, structure, and metadata across supported file formats."""

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".jpeg",
        ".jpg",
        ".png",
    }

    @classmethod
    def extract_file(cls, file_path: str) -> DocumentExtractionResult:
        """Extract content and structure from a file path."""
        if not os.path.exists(file_path):
            return DocumentExtractionResult(
                filename=os.path.basename(file_path),
                file_path=file_path,
                file_type="",
                file_size_bytes=0,
                pages=[],
                status="error",
                error_message="File does not exist on disk",
            )

        filename = os.path.basename(file_path)
        ext = os.path.splitext(filename)[1].lower()
        file_size = os.path.getsize(file_path)

        if ext == ".pdf":
            return cls._extract_pdf(file_path, filename, ext, file_size)
        elif ext == ".docx":
            return cls._extract_docx(file_path, filename, ext, file_size)
        elif ext in (".txt", ".md"):
            return cls._extract_text(file_path, filename, ext, file_size)
        elif ext == ".csv":
            return cls._extract_csv(file_path, filename, ext, file_size)
        elif ext == ".json":
            return cls._extract_json(file_path, filename, ext, file_size)
        elif ext in (".jpeg", ".jpg", ".png"):
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="scanned_ocr_required",
                error_message="Image file requires Optical Character Recognition (OCR)",
            )
        else:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="unsupported",
                error_message=f"Unsupported file format: {ext}",
            )

    @classmethod
    def _extract_pdf(
        cls, file_path: str, filename: str, ext: str, file_size: int
    ) -> DocumentExtractionResult:
        """Extract text page-by-page from a PDF, detecting scanned/image documents."""
        try:
            reader = PdfReader(file_path)
            extracted_pages: List[ExtractedPage] = []
            non_empty_count = 0

            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                clean_text = text.strip()
                is_empty = len(clean_text) == 0

                if not is_empty:
                    non_empty_count += 1

                # Extract first meaningful line as candidate heading
                first_line = clean_text.split("\n")[0] if clean_text else None
                heading = (
                    first_line[:120].strip()
                    if first_line and len(first_line) > 3
                    else None
                )

                extracted_pages.append(
                    ExtractedPage(
                        page_number=idx + 1,
                        text=clean_text,
                        section_heading=heading,
                        is_empty=is_empty,
                    )
                )

            # If all pages are completely empty, this is an image-based/scanned PDF
            if non_empty_count == 0 and len(extracted_pages) > 0:
                return DocumentExtractionResult(
                    filename=filename,
                    file_path=file_path,
                    file_type=ext,
                    file_size_bytes=file_size,
                    pages=extracted_pages,
                    status="scanned_ocr_required",
                    error_message="Scanned or image-based PDF. Text layer not present, requires OCR.",
                    metadata={"page_count": len(extracted_pages)},
                )

            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=extracted_pages,
                status="success",
                metadata={"page_count": len(extracted_pages)},
            )
        except Exception as e:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="error",
                error_message=f"PDF extraction failed: {str(e)}",
            )

    @classmethod
    def _extract_docx(
        cls, file_path: str, filename: str, ext: str, file_size: int
    ) -> DocumentExtractionResult:
        """Extract structured paragraphs, headings, and tables from a DOCX document."""
        try:
            doc = docx.Document(file_path)
            extracted_pages: List[ExtractedPage] = []

            current_heading: Optional[str] = None
            current_paragraphs: List[str] = []
            current_tables: List[List[str]] = []
            section_idx = 1

            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue

                # Detect heading styles
                if p.style and (
                    "heading" in p.style.name.lower() or "title" in p.style.name.lower()
                ):
                    if current_paragraphs:
                        extracted_pages.append(
                            ExtractedPage(
                                page_number=section_idx,
                                text="\n\n".join(current_paragraphs),
                                section_heading=current_heading,
                                tables=current_tables,
                            )
                        )
                        section_idx += 1
                        current_paragraphs = []
                        current_tables = []
                    current_heading = txt
                else:
                    current_paragraphs.append(txt)

            # Extract tables
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(row_cells):
                        table_rows.append(" | ".join(row_cells))
                if table_rows:
                    current_paragraphs.append(
                        "\n[Table Data]\n" + "\n".join(table_rows)
                    )
                    current_tables.append(table_rows)

            if current_paragraphs:
                extracted_pages.append(
                    ExtractedPage(
                        page_number=section_idx,
                        text="\n\n".join(current_paragraphs),
                        section_heading=current_heading,
                        tables=current_tables,
                    )
                )

            if not extracted_pages:
                extracted_pages.append(
                    ExtractedPage(page_number=1, text="", is_empty=True)
                )

            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=extracted_pages,
                status="success",
                metadata={"section_count": len(extracted_pages)},
            )
        except Exception as e:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="error",
                error_message=f"DOCX extraction failed: {str(e)}",
            )

    @classmethod
    def _extract_text(
        cls, file_path: str, filename: str, ext: str, file_size: int
    ) -> DocumentExtractionResult:
        """Extract plain text or markdown files."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()

            clean_text = content.strip()
            pages = [
                ExtractedPage(
                    page_number=1,
                    text=clean_text,
                    is_empty=len(clean_text) == 0,
                )
            ]
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=pages,
                status="success",
            )
        except Exception as e:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="error",
                error_message=f"Text extraction failed: {str(e)}",
            )

    @classmethod
    def _extract_csv(
        cls, file_path: str, filename: str, ext: str, file_size: int
    ) -> DocumentExtractionResult:
        """Extract CSV tabular facts."""
        try:
            rows = []
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    if any(row):
                        rows.append(" | ".join([c.strip() for c in row]))

            text = "\n".join(rows)
            pages = [
                ExtractedPage(
                    page_number=1,
                    text=text,
                    is_empty=len(text) == 0,
                )
            ]
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=pages,
                status="success",
            )
        except Exception as e:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="error",
                error_message=f"CSV extraction failed: {str(e)}",
            )

    @classmethod
    def _extract_json(
        cls, file_path: str, filename: str, ext: str, file_size: int
    ) -> DocumentExtractionResult:
        """Extract JSON structured data."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json.load(f)

            text = json.dumps(data, indent=2, ensure_ascii=False)
            pages = [
                ExtractedPage(
                    page_number=1,
                    text=text,
                    is_empty=len(text.strip()) == 0,
                )
            ]
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=pages,
                status="success",
            )
        except Exception as e:
            return DocumentExtractionResult(
                filename=filename,
                file_path=file_path,
                file_type=ext,
                file_size_bytes=file_size,
                pages=[],
                status="error",
                error_message=f"JSON extraction failed: {str(e)}",
            )
